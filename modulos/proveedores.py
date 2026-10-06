"""Módulo Proveedores: CRUD completo, filtros en tiempo real, detalle y exportación.

Estructura (POO):
    ValidationError        -> errores de validación de formulario
    ProveedorValidator     -> reglas de RUC/DNI, correo, teléfono, longitudes
    ProveedorRepository    -> todas las consultas SQL del módulo
    ProveedorExporter      -> exportación a Excel (.xlsx) y PDF
    ProveedorFormDialog    -> ventana Nuevo / Editar
    ProveedorDetalleDialog -> ficha + historial de compras + productos suministrados
    ProveedoresView        -> pantalla principal (tabla, filtros, acciones)
"""

from __future__ import annotations

import re
import tkinter as tk
from datetime import datetime
from decimal import Decimal
from tkinter import filedialog, messagebox, ttk

from mysql.connector import Error as MySQLError

import database
from .comun import abrir_archivo, escape_like, mensaje_error_mysql, registrar_log

# Si es False solo se valida longitud/prefijo del RUC, no el dígito verificador.
VALIDAR_DIGITO_RUC = True

ESTADOS = {"A": "Activo", "I": "Inactivo"}
ESTADO_A_CODIGO = {v: k for k, v in ESTADOS.items()}
BG = "#eef2f7"


# =============================================================
# Validación
# =============================================================
class ValidationError(Exception):
    def __init__(self, errores):
        self.errores = list(errores)
        super().__init__("\n".join(self.errores))


class ProveedorValidator:
    EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}$")
    TEL_RE = re.compile(r"^\+?[0-9]{6,15}$")
    LIMITES = {"razon_social": 100, "contacto": 100, "correo": 80, "telefono": 20, "direccion": 150}

    @staticmethod
    def ruc_digito_ok(ruc: str) -> bool:
        """Dígito verificador del RUC peruano (módulo 11)."""
        pesos = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
        suma = sum(int(d) * p for d, p in zip(ruc[:10], pesos))
        return (11 - suma % 11) % 10 == int(ruc[10])

    @classmethod
    def error_documento(cls, doc: str):
        if not doc:
            return "El RUC / DNI es obligatorio."
        if not re.fullmatch(r"[0-9]+", doc):
            return "El RUC / DNI solo debe contener números."
        if len(doc) == 8:
            return None  # DNI
        if len(doc) == 11:
            if doc[:2] not in ("10", "15", "16", "17", "20"):
                return "El RUC debe empezar con 10, 15, 16, 17 o 20."
            if VALIDAR_DIGITO_RUC and not cls.ruc_digito_ok(doc):
                return "El RUC no es válido (dígito verificador incorrecto)."
            return None
        return "Ingresa un DNI (8 dígitos) o un RUC (11 dígitos)."

    @classmethod
    def validar(cls, datos: dict) -> dict:
        """Devuelve los datos limpios o lanza ValidationError con TODOS los problemas."""
        errores = []
        limpio = {k: (datos.get(k) or "").strip() for k in
                  ("ruc", "razon_social", "contacto", "correo", "telefono", "direccion")}
        limpio["estado"] = datos.get("estado") if datos.get("estado") in ESTADOS else "A"

        msg = cls.error_documento(limpio["ruc"])
        if msg:
            errores.append(msg)

        if len(limpio["razon_social"]) < 3:
            errores.append("La razón social / nombre comercial es obligatoria (mínimo 3 caracteres).")

        if limpio["correo"]:
            limpio["correo"] = limpio["correo"].lower()
            if not cls.EMAIL_RE.match(limpio["correo"]):
                errores.append("El correo electrónico no tiene un formato válido (ej.: ventas@empresa.com).")

        if limpio["telefono"]:
            tel = re.sub(r"[\s\-()]", "", limpio["telefono"])
            if not cls.TEL_RE.match(tel):
                errores.append("El teléfono debe ser numérico (6 a 15 dígitos, puede iniciar con +).")
            else:
                limpio["telefono"] = tel

        for campo, maximo in cls.LIMITES.items():
            if len(limpio[campo]) > maximo:
                errores.append(f"«{campo.replace('_', ' ')}» no puede superar {maximo} caracteres.")

        if errores:
            raise ValidationError(errores)
        return limpio


# =============================================================
# Acceso a datos
# =============================================================
class ProveedorRepository:
    COLUMNAS = "id_proveedor, ruc, razon_social, contacto, correo, telefono, direccion, estado"

    def listar(self, termino: str = "", estado: str | None = None):
        sql = f"SELECT {self.COLUMNAS} FROM proveedores WHERE 1=1"
        params: list = []
        if termino:
            like = f"%{escape_like(termino)}%"
            sql += " AND (ruc LIKE %s OR razon_social LIKE %s OR telefono LIKE %s OR contacto LIKE %s)"
            params += [like] * 4
        if estado in ESTADOS:
            sql += " AND estado=%s"
            params.append(estado)
        sql += " ORDER BY razon_social ASC"
        return database.fetch_all(sql, tuple(params))

    def obtener(self, id_proveedor: int):
        return database.fetch_one("SELECT * FROM proveedores WHERE id_proveedor=%s", (id_proveedor,))

    def existe_documento(self, ruc: str, excluir_id: int | None = None) -> bool:
        sql = "SELECT id_proveedor FROM proveedores WHERE ruc=%s"
        params = [ruc]
        if excluir_id:
            sql += " AND id_proveedor<>%s"
            params.append(excluir_id)
        return database.fetch_one(sql, tuple(params)) is not None

    def crear(self, d: dict) -> int:
        return database.execute(
            "INSERT INTO proveedores(ruc,razon_social,contacto,correo,telefono,direccion,estado) "
            "VALUES(%s,%s,%s,%s,%s,%s,%s)",
            (d["ruc"], d["razon_social"], d["contacto"] or None, d["correo"] or None,
             d["telefono"] or None, d["direccion"] or None, d["estado"]),
        )

    def actualizar(self, id_proveedor: int, d: dict) -> None:
        database.execute(
            "UPDATE proveedores SET ruc=%s, razon_social=%s, contacto=%s, correo=%s, "
            "telefono=%s, direccion=%s, estado=%s WHERE id_proveedor=%s",
            (d["ruc"], d["razon_social"], d["contacto"] or None, d["correo"] or None,
             d["telefono"] or None, d["direccion"] or None, d["estado"], id_proveedor),
        )

    def cambiar_estado(self, id_proveedor: int, estado: str) -> None:
        database.execute("UPDATE proveedores SET estado=%s WHERE id_proveedor=%s", (estado, id_proveedor))

    def contar_ingresos(self, id_proveedor: int) -> int:
        row = database.fetch_one(
            "SELECT COUNT(*) AS n FROM ingreso_mercaderia WHERE id_proveedor=%s", (id_proveedor,))
        return int(row["n"]) if row else 0

    def eliminar(self, id_proveedor: int) -> None:
        database.execute("DELETE FROM proveedores WHERE id_proveedor=%s", (id_proveedor,))

    def historial_compras(self, id_proveedor: int):
        return database.fetch_all(
            "SELECT id_ingreso, fecha_ingreso, total_ingreso FROM ingreso_mercaderia "
            "WHERE id_proveedor=%s ORDER BY fecha_ingreso DESC, id_ingreso DESC", (id_proveedor,))

    def productos_suministrados(self, id_proveedor: int):
        return database.fetch_all(
            "SELECT p.nombre_producto, p.marca, SUM(d.cantidad) AS cantidad, "
            "AVG(d.costo_unitario) AS costo_promedio, MAX(i.fecha_ingreso) AS ultima_fecha "
            "FROM detalle_ingreso d "
            "JOIN ingreso_mercaderia i ON i.id_ingreso=d.id_ingreso "
            "JOIN productos p ON p.id_producto=d.id_producto "
            "WHERE i.id_proveedor=%s "
            "GROUP BY p.id_producto, p.nombre_producto, p.marca "
            "ORDER BY ultima_fecha DESC", (id_proveedor,))


# =============================================================
# Exportación
# =============================================================
class ProveedorExporter:
    CABECERAS = ["ID", "RUC / DNI", "Razón social", "Contacto", "Correo", "Teléfono", "Dirección", "Estado"]

    @staticmethod
    def _filas(proveedores):
        return [[p["id_proveedor"], p["ruc"] or "", p["razon_social"] or "", p["contacto"] or "",
                 p["correo"] or "", p["telefono"] or "", p["direccion"] or "",
                 ESTADOS.get(p["estado"], p["estado"])] for p in proveedores]

    def a_excel(self, ruta: str, proveedores) -> None:
        from openpyxl import Workbook                       # pip install openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Proveedores"
        ws.append(self.CABECERAS)
        for fila in self._filas(proveedores):
            ws.append(fila)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="243B53")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, ancho in enumerate((7, 16, 38, 26, 30, 16, 42, 11), start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        for fila in ws.iter_rows(min_row=2, min_col=2, max_col=2):   # RUC/DNI como texto
            fila[0].number_format = "@"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(ruta)

    def a_pdf(self, ruta: str, proveedores) -> None:
        from reportlab.lib import colors                    # pip install reportlab
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        estilos = getSampleStyleSheet()
        celda = ParagraphStyle("celda", parent=estilos["Normal"], fontSize=8, leading=10)
        cab = ParagraphStyle("cab", parent=celda, textColor=colors.white, fontName="Helvetica-Bold")

        datos = [[Paragraph(h, cab) for h in self.CABECERAS]]
        for fila in self._filas(proveedores):
            datos.append([Paragraph(str(v), celda) for v in fila])

        tabla = Table(datos, repeatRows=1,
                      colWidths=[10 * mm, 25 * mm, 55 * mm, 40 * mm, 48 * mm, 25 * mm, 55 * mm, 18 * mm])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B53")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BCCCDC")),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))

        doc = SimpleDocTemplate(ruta, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                                topMargin=14 * mm, bottomMargin=14 * mm, title="Listado de proveedores")
        doc.build([
            Paragraph("Listado de proveedores", estilos["Title"]),
            Paragraph(f"Generado el {datetime.now():%d/%m/%Y %H:%M} — {len(proveedores)} registro(s)",
                      estilos["Normal"]),
            Spacer(1, 6 * mm), tabla,
        ])


# =============================================================
# Ventana Nuevo / Editar
# =============================================================
class ProveedorFormDialog(tk.Toplevel):
    def __init__(self, parent, repo: ProveedorRepository, user: dict,
                 id_proveedor: int | None = None, on_saved=None):
        super().__init__(parent)
        self.repo, self.user, self.id = repo, user, id_proveedor
        self.on_saved = on_saved
        self.title("Editar proveedor" if id_proveedor else "Nuevo proveedor")
        self.resizable(False, False)
        self.transient(parent)

        self.vars = {k: tk.StringVar() for k in
                     ("ruc", "razon_social", "contacto", "correo", "telefono", "direccion")}
        self.estado_var = tk.StringVar(value="Activo")
        if id_proveedor:
            fila = repo.obtener(id_proveedor)
            if fila:
                for k, v in self.vars.items():
                    v.set(str(fila.get(k) or ""))
                self.estado_var.set(ESTADOS.get(fila.get("estado"), "Activo"))

        self._construir()
        self.bind("<Return>", lambda _e: self._guardar())
        self.bind("<Escape>", lambda _e: self.destroy())
        self.update_idletasks()
        x = parent.winfo_rootx() + (parent.winfo_width() - self.winfo_width()) // 2
        y = parent.winfo_rooty() + (parent.winfo_height() - self.winfo_height()) // 3
        self.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        self.grab_set()
        self.entry_ruc.focus_set()

    def _construir(self):
        marco = tk.Frame(self, padx=26, pady=20)
        marco.pack(fill="both", expand=True)
        marco.columnconfigure(1, weight=1)

        solo_doc = (self.register(lambda s: re.fullmatch(r"[0-9]{0,11}", s) is not None), "%P")
        solo_tel = (self.register(lambda s: re.fullmatch(r"\+?[0-9 \-()]{0,20}", s) is not None), "%P")

        campos = [
            ("RUC / DNI *", "ruc", solo_doc),
            ("Razón social / Nombre comercial *", "razon_social", None),
            ("Contacto / Vendedor", "contacto", None),
            ("Correo electrónico", "correo", None),
            ("Teléfono", "telefono", solo_tel),
            ("Dirección", "direccion", None),
        ]
        for fila, (texto, clave, vcmd) in enumerate(campos):
            tk.Label(marco, text=texto, font=("Segoe UI", 10, "bold"), fg="#334e68",
                     anchor="w").grid(row=fila, column=0, sticky="w", pady=7, padx=(0, 14))
            kw = {"validate": "key", "validatecommand": vcmd} if vcmd else {}
            e = ttk.Entry(marco, textvariable=self.vars[clave], width=44, **kw)
            e.grid(row=fila, column=1, sticky="ew", pady=7)
            if clave == "ruc":
                self.entry_ruc = e

        tk.Label(marco, text="Estado", font=("Segoe UI", 10, "bold"), fg="#334e68").grid(
            row=len(campos), column=0, sticky="w", pady=7)
        ttk.Combobox(marco, textvariable=self.estado_var, values=list(ESTADOS.values()),
                     state="readonly", width=14).grid(row=len(campos), column=1, sticky="w", pady=7)

        tk.Label(marco, text="* Campos obligatorios", fg="#7b8794", font=("Segoe UI", 9)).grid(
            row=len(campos) + 1, column=0, columnspan=2, sticky="w", pady=(6, 0))

        botones = tk.Frame(marco)
        botones.grid(row=len(campos) + 2, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(botones, text="Cancelar", command=self.destroy).pack(side="right")
        ttk.Button(botones, text="Guardar", style="Accent.TButton",
                   command=self._guardar).pack(side="right", padx=8)

    def _guardar(self):
        datos = {k: v.get() for k, v in self.vars.items()}
        datos["estado"] = ESTADO_A_CODIGO.get(self.estado_var.get(), "A")
        try:
            limpio = ProveedorValidator.validar(datos)
            if self.repo.existe_documento(limpio["ruc"], self.id):
                raise ValidationError([f"Ya existe un proveedor registrado con el RUC / DNI {limpio['ruc']}."])
            if self.id:
                self.repo.actualizar(self.id, limpio)
                rid, accion = self.id, "Actualizó proveedor"
            else:
                rid, accion = self.repo.crear(limpio), "Registró proveedor"
            registrar_log(self.user["id_usuario"], accion, "proveedores", rid)
        except ValidationError as exc:
            messagebox.showwarning("Revisa los datos",
                                   "\n".join(f"• {m}" for m in exc.errores), parent=self)
            return
        except MySQLError as exc:
            if getattr(exc, "errno", None) == 1062:
                msg = "Ya existe un proveedor con ese RUC / DNI."
            else:
                msg = mensaje_error_mysql(exc)
            messagebox.showerror("No se pudo guardar", msg, parent=self)
            return

        messagebox.showinfo("Proveedor", "Proveedor guardado correctamente.", parent=self)
        if self.on_saved:
            self.on_saved()
        self.destroy()


# =============================================================
# Ventana Ver detalles
# =============================================================
class ProveedorDetalleDialog(tk.Toplevel):
    def __init__(self, parent, repo: ProveedorRepository, id_proveedor: int):
        super().__init__(parent)
        self.title("Detalle del proveedor")
        self.geometry("820x560")
        self.transient(parent)
        self.bind("<Escape>", lambda _e: self.destroy())

        try:
            p = repo.obtener(id_proveedor)
            compras = repo.historial_compras(id_proveedor)
            try:
                productos = repo.productos_suministrados(id_proveedor)
            except MySQLError:       # detalle_ingreso todavía no creada
                productos = []
        except MySQLError as exc:
            messagebox.showerror("Detalle", mensaje_error_mysql(exc), parent=parent)
            self.destroy()
            return
        if not p:
            messagebox.showwarning("Detalle", "El proveedor ya no existe.", parent=parent)
            self.destroy()
            return

        tk.Label(self, text=p["razon_social"], font=("Segoe UI", 16, "bold"), fg="#243b53",
                 anchor="w").pack(fill="x", padx=22, pady=(18, 0))
        tk.Label(self, text=f"{ESTADOS.get(p['estado'], '')}  ·  RUC / DNI: {p['ruc'] or 'pendiente'}",
                 font=("Segoe UI", 10), fg="#7b8794", anchor="w").pack(fill="x", padx=22, pady=(0, 10))

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=18, pady=(0, 10))

        # --- Ficha ---
        ficha = tk.Frame(nb, padx=24, pady=18)
        nb.add(ficha, text="Ficha técnica")
        registrado = p.get("fecha_registro")
        datos = [("RUC / DNI", p["ruc"] or "— pendiente —"), ("Razón social", p["razon_social"]),
                 ("Contacto / Vendedor", p["contacto"]), ("Correo", p["correo"]),
                 ("Teléfono", p["telefono"]), ("Dirección", p["direccion"]),
                 ("Estado", ESTADOS.get(p["estado"], p["estado"])),
                 ("Registrado", registrado.strftime("%d/%m/%Y %H:%M") if registrado else "—")]
        ficha.columnconfigure(1, weight=1)
        for i, (k, v) in enumerate(datos):
            tk.Label(ficha, text=k, font=("Segoe UI", 10, "bold"), fg="#334e68").grid(
                row=i, column=0, sticky="nw", pady=6, padx=(0, 20))
            tk.Label(ficha, text=str(v or "—"), font=("Segoe UI", 10), anchor="w",
                     justify="left", wraplength=520).grid(row=i, column=1, sticky="w", pady=6)

        # --- Historial de compras ---
        hist = tk.Frame(nb, padx=12, pady=12)
        nb.add(hist, text=f"Historial de compras ({len(compras)})")
        total = sum((Decimal(str(c["total_ingreso"] or 0)) for c in compras), Decimal("0"))
        tk.Label(hist, text=f"{len(compras)} compra(s)  ·  Total comprado: S/ {total:,.2f}",
                 font=("Segoe UI", 10, "bold"), fg="#334e68").pack(anchor="w", pady=(0, 8))
        t1 = self._tabla(hist, ("n", "fecha", "total"), ("N° ingreso", "Fecha", "Total"), (110, 140, 140))
        for c in compras:
            t1.insert("", "end", values=(c["id_ingreso"], c["fecha_ingreso"],
                                         f"S/ {Decimal(str(c['total_ingreso'] or 0)):,.2f}"))
        if not compras:
            tk.Label(hist, text="Este proveedor aún no tiene compras registradas.",
                     fg="#7b8794").pack(pady=6)

        # --- Productos suministrados ---
        prod = tk.Frame(nb, padx=12, pady=12)
        nb.add(prod, text=f"Productos suministrados ({len(productos)})")
        t2 = self._tabla(prod, ("producto", "marca", "cant", "costo", "ultima"),
                         ("Producto", "Marca", "Cant. total", "Costo prom.", "Última compra"),
                         (280, 100, 90, 100, 110))
        for r in productos:
            t2.insert("", "end", values=(r["nombre_producto"], r["marca"] or "", int(r["cantidad"] or 0),
                                         f"S/ {Decimal(str(r['costo_promedio'] or 0)):,.2f}",
                                         r["ultima_fecha"]))
        if not productos:
            tk.Label(prod, text="Aún no hay productos asociados a este proveedor.",
                     fg="#7b8794").pack(pady=6)

        ttk.Button(self, text="Cerrar", command=self.destroy).pack(side="right", padx=22, pady=(0, 14))

    @staticmethod
    def _tabla(parent, cols, titulos, anchos):
        marco = tk.Frame(parent)
        marco.pack(fill="both", expand=True)
        tree = ttk.Treeview(marco, columns=cols, show="headings")
        for c, t, w in zip(cols, titulos, anchos):
            tree.heading(c, text=t)
            tree.column(c, width=w, anchor="w")
        sb = ttk.Scrollbar(marco, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        return tree


# =============================================================
# Pantalla principal
# =============================================================
class ProveedoresView(tk.Frame):
    COLUMNAS = (("id", "ID", 50), ("ruc", "RUC / DNI", 105), ("razon", "Razón social", 230),
                ("contacto", "Contacto", 150), ("telefono", "Teléfono", 105),
                ("correo", "Correo", 190), ("direccion", "Dirección", 210), ("estado", "Estado", 80))

    def __init__(self, parent, root, user: dict):
        super().__init__(parent, bg=BG)
        self.root, self.user = root, user
        self.repo = ProveedorRepository()
        self.exporter = ProveedorExporter()
        self._filas: list = []
        self._after_id = None
        self._orden_desc: dict = {}
        self.busqueda = tk.StringVar()
        self.filtro_estado = tk.StringVar(value="Todos")

        self._construir()
        self.cargar()

    # ---------- construcción ----------
    def _construir(self):
        barra = tk.Frame(self, bg=BG)
        barra.pack(fill="x", padx=28, pady=(4, 4))
        tk.Label(barra, text="Buscar", bg=BG, fg="#334e68", font=("Segoe UI", 10, "bold")).pack(side="left")
        entrada = ttk.Entry(barra, textvariable=self.busqueda, width=34)
        entrada.pack(side="left", padx=(8, 4))
        ttk.Button(barra, text="✖", width=3, command=lambda: self.busqueda.set("")).pack(side="left")
        tk.Label(barra, text="Estado", bg=BG, fg="#334e68",
                 font=("Segoe UI", 10, "bold")).pack(side="left", padx=(18, 6))
        combo = ttk.Combobox(barra, textvariable=self.filtro_estado, state="readonly", width=14,
                             values=["Todos", "Solo activos", "Solo inactivos"])
        combo.pack(side="left")
        self.lbl_total = tk.Label(barra, text="", bg=BG, fg="#7b8794", font=("Segoe UI", 9))
        self.lbl_total.pack(side="right")
        tk.Label(self, text="Busca por RUC / DNI, razón social, teléfono o contacto — "
                            "los resultados se actualizan mientras escribes.",
                 bg=BG, fg="#9aa5b1", font=("Segoe UI", 8)).pack(anchor="w", padx=28)

        acciones = tk.Frame(self, bg=BG)
        acciones.pack(fill="x", padx=28, pady=(8, 4))
        ttk.Button(acciones, text="＋ Nuevo proveedor", style="Accent.TButton",
                   command=self.nuevo).pack(side="left")
        self.btn_editar = ttk.Button(acciones, text="✏️ Editar", command=self.editar)
        self.btn_eliminar = ttk.Button(acciones, text="🗑️ Eliminar", command=self.eliminar)
        self.btn_detalle = ttk.Button(acciones, text="👁️ Ver detalles", command=self.ver_detalle)
        for b in (self.btn_editar, self.btn_eliminar, self.btn_detalle):
            b.pack(side="left", padx=(8, 0))
        exportar = ttk.Menubutton(acciones, text="📄 Exportar ▾")
        menu_exp = tk.Menu(exportar, tearoff=0)
        menu_exp.add_command(label="Excel (.xlsx)", command=lambda: self.exportar("xlsx"))
        menu_exp.add_command(label="PDF (.pdf)", command=lambda: self.exportar("pdf"))
        exportar["menu"] = menu_exp
        exportar.pack(side="right")

        marco = tk.Frame(self, bg="white")
        marco.pack(fill="both", expand=True, padx=28, pady=(4, 12))
        cols = [c[0] for c in self.COLUMNAS]
        self.tree = ttk.Treeview(marco, columns=cols, show="headings", selectmode="browse")
        for clave, titulo, ancho in self.COLUMNAS:
            self.tree.heading(clave, text=titulo, command=lambda c=clave: self._ordenar(c))
            self.tree.column(clave, width=ancho, anchor="center" if clave in ("id", "estado") else "w")
        sb_y = ttk.Scrollbar(marco, orient="vertical", command=self.tree.yview)
        sb_x = ttk.Scrollbar(marco, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=sb_y.set, xscrollcommand=sb_x.set)
        sb_y.pack(side="right", fill="y")
        sb_x.pack(side="bottom", fill="x")
        self.tree.pack(fill="both", expand=True, padx=(10, 0), pady=(10, 0))
        self.tree.tag_configure("inactivo", foreground="#9aa5b1")
        self.tree.tag_configure("par", background="#f7fafc")

        self.menu = tk.Menu(self, tearoff=0)
        self.menu.add_command(label="✏️  Editar", command=self.editar)
        self.menu.add_command(label="👁️  Ver detalles", command=self.ver_detalle)
        self.menu.add_separator()
        self.menu.add_command(label="🗑️  Eliminar", command=self.eliminar)

        self.tree.bind("<<TreeviewSelect>>", lambda _e: self._actualizar_botones())
        self.tree.bind("<Double-1>", lambda _e: self.editar())
        self.tree.bind("<Button-3>", self._menu_contextual)
        self.busqueda.trace_add("write", lambda *_: self._programar_carga())
        combo.bind("<<ComboboxSelected>>", lambda _e: self.cargar())
        self._actualizar_botones()
        entrada.focus_set()

    # ---------- datos ----------
    def _programar_carga(self):
        """Espera 250 ms tras la última tecla para no consultar MySQL por cada letra."""
        if self._after_id:
            self.after_cancel(self._after_id)
        self._after_id = self.after(250, self.cargar)

    def cargar(self, seleccionar_id: int | None = None):
        self._after_id = None
        estado = {"Solo activos": "A", "Solo inactivos": "I"}.get(self.filtro_estado.get())
        try:
            self._filas = self.repo.listar(self.busqueda.get().strip(), estado)
        except MySQLError as exc:
            messagebox.showerror("Proveedores", mensaje_error_mysql(exc), parent=self.root)
            return
        self.tree.delete(*self.tree.get_children())
        for i, p in enumerate(self._filas):
            tags = ["inactivo"] if p["estado"] == "I" else (["par"] if i % 2 else [])
            self.tree.insert("", "end", iid=str(p["id_proveedor"]), tags=tags, values=(
                p["id_proveedor"], p["ruc"] or "— pendiente —", p["razon_social"] or "",
                p["contacto"] or "", p["telefono"] or "", p["correo"] or "",
                p["direccion"] or "", ESTADOS.get(p["estado"], p["estado"])))
        self.lbl_total.config(text=f"{len(self._filas)} proveedor(es)")
        if seleccionar_id and self.tree.exists(str(seleccionar_id)):
            self.tree.selection_set(str(seleccionar_id))
        self._actualizar_botones()

    def _seleccionado(self):
        sel = self.tree.selection()
        if not sel:
            return None
        id_sel = int(sel[0])
        return next((p for p in self._filas if p["id_proveedor"] == id_sel), None)

    def _actualizar_botones(self):
        estado = "normal" if self.tree.selection() else "disabled"
        for b in (self.btn_editar, self.btn_eliminar, self.btn_detalle):
            b.configure(state=estado)

    def _menu_contextual(self, evento):
        fila = self.tree.identify_row(evento.y)
        if fila:
            self.tree.selection_set(fila)
            self.menu.tk_popup(evento.x_root, evento.y_root)

    def _ordenar(self, columna):
        desc = self._orden_desc.get(columna, False)
        items = [(self.tree.set(i, columna), i) for i in self.tree.get_children("")]
        clave = (lambda t: int(t[0])) if columna == "id" else (lambda t: t[0].lower())
        items.sort(key=clave, reverse=desc)
        for pos, (_, iid) in enumerate(items):
            self.tree.move(iid, "", pos)
        self._orden_desc[columna] = not desc

    # ---------- acciones ----------
    def nuevo(self):
        ProveedorFormDialog(self.root, self.repo, self.user, on_saved=self.cargar)

    def editar(self):
        p = self._seleccionado()
        if p:
            ProveedorFormDialog(self.root, self.repo, self.user, p["id_proveedor"],
                                on_saved=lambda: self.cargar(p["id_proveedor"]))

    def ver_detalle(self):
        p = self._seleccionado()
        if p:
            ProveedorDetalleDialog(self.root, self.repo, p["id_proveedor"])

    def eliminar(self):
        p = self._seleccionado()
        if not p:
            return
        nombre, pid = p["razon_social"], p["id_proveedor"]
        try:
            if p["estado"] == "A":
                if messagebox.askyesno(
                        "Confirmar eliminación",
                        f"¿Deseas eliminar al proveedor «{nombre}»?\n\n"
                        "Se marcará como INACTIVO (eliminación lógica): conservará su historial "
                        "de compras y podrás reactivarlo desde Editar.", parent=self.root):
                    self.repo.cambiar_estado(pid, "I")
                    registrar_log(self.user["id_usuario"], "Desactivó proveedor", "proveedores", pid)
                    messagebox.showinfo("Proveedor", "Proveedor marcado como inactivo.", parent=self.root)
                    self.cargar()
                return

            # Ya estaba inactivo: solo se puede borrar de forma definitiva si no tiene compras.
            compras = self.repo.contar_ingresos(pid)
            if compras:
                messagebox.showinfo(
                    "No se puede eliminar definitivamente",
                    f"«{nombre}» tiene {compras} compra(s) registradas, por eso se conserva como inactivo.\n"
                    "Puedes reactivarlo desde Editar.", parent=self.root)
            elif messagebox.askyesno(
                    "Eliminar definitivamente",
                    f"«{nombre}» ya está inactivo y no tiene compras.\n\n"
                    "¿Eliminarlo DEFINITIVAMENTE? Esta acción no se puede deshacer.",
                    icon="warning", parent=self.root):
                self.repo.eliminar(pid)
                registrar_log(self.user["id_usuario"], "Eliminó proveedor", "proveedores", pid)
                messagebox.showinfo("Proveedor", "Proveedor eliminado.", parent=self.root)
                self.cargar()
        except MySQLError as exc:
            messagebox.showerror("No se pudo eliminar", mensaje_error_mysql(exc), parent=self.root)

    def exportar(self, formato: str):
        if not self._filas:
            messagebox.showinfo("Exportar", "No hay proveedores para exportar con el filtro actual.",
                                parent=self.root)
            return
        tipos = {"xlsx": ("Libro de Excel", "*.xlsx"), "pdf": ("Documento PDF", "*.pdf")}
        ruta = filedialog.asksaveasfilename(
            parent=self.root, title="Exportar proveedores", defaultextension=f".{formato}",
            filetypes=[tipos[formato]], initialfile=f"Proveedores_{datetime.now():%Y%m%d_%H%M}")
        if not ruta:
            return
        try:
            (self.exporter.a_excel if formato == "xlsx" else self.exporter.a_pdf)(ruta, self._filas)
        except ImportError:
            libreria = "openpyxl" if formato == "xlsx" else "reportlab"
            messagebox.showerror("Exportar", f"Falta instalar la librería.\n\npip install {libreria}",
                                 parent=self.root)
            return
        except OSError as exc:
            messagebox.showerror("Exportar", f"No se pudo escribir el archivo.\n¿Está abierto en otro programa?\n\n{exc}",
                                 parent=self.root)
            return
        if messagebox.askyesno("Exportar", f"Archivo creado:\n{ruta}\n\n¿Deseas abrirlo ahora?", parent=self.root):
            abrir_archivo(ruta)
